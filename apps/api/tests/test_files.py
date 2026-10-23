import io

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.templates import finance_workbook
from app.services.xlsx import snapshot_to_csv, snapshot_to_xlsx, xlsx_to_snapshot
from tests.conftest import personal_workspace, register


def _sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws["A1"] = "Товар"
    ws["A1"].font = Font(bold=True, color="FF0000FF")
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    ws["B1"] = 10
    ws["B2"] = 32.5
    ws["B3"] = "=SUM(B1:B2)"
    ws["C1"] = "Статус"
    ws.merge_cells("D1:E1")
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"Да,Нет"')
    dv.add("C2:C10")
    ws.add_data_validation(dv)
    wb.create_sheet("Второй")["A1"] = "=Данные!B3*2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_xlsx_import_maps_values_styles_and_layout() -> None:
    snapshot = xlsx_to_snapshot(_sample_xlsx(), "Тест")
    first, second = (snapshot["sheets"][s] for s in snapshot["sheetOrder"])
    assert first["name"] == "Данные" and second["name"] == "Второй"
    cells = first["cellData"]
    assert cells["0"]["0"]["v"] == "Товар"
    style = snapshot["styles"][cells["0"]["0"]["s"]]
    assert style["bl"] == 1 and style["cl"] == {"rgb": "#0000FF"}
    assert style["bg"] == {"rgb": "#FFFF00"}
    assert cells["2"]["1"]["f"] == "=SUM(B1:B2)"
    assert first["mergeData"] == [{"startRow": 0, "endRow": 0, "startColumn": 3, "endColumn": 4}]
    assert first["columnData"]["0"]["w"] == 215
    assert first["freeze"]["ySplit"] == 1
    assert second["cellData"]["0"]["0"]["f"] == "=Данные!B3*2"
    assert snapshot["resources"][0]["name"] == "SHEET_DATA_VALIDATION_PLUGIN"


def test_finance_template_roundtrips_to_xlsx() -> None:
    snapshot = finance_workbook("Финансы")
    data = snapshot_to_xlsx(snapshot)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Обзор", "Октябрь", "Категории"]
    overview = wb["Обзор"]
    assert overview["A1"].value == "Категория"
    assert overview["C3"].value.startswith("=SUMIF(Октябрь!")
    assert overview["B9"].value == "=SUM(B3:B8)"
    assert overview["B2"].number_format == '#,##0 "₽"'
    assert overview["A1"].font.b
    assert overview.freeze_panes == "A2"
    assert overview.data_validations.dataValidation[0].formula1 == '"Выполнено,В процессе,Ожидает"'

    back = xlsx_to_snapshot(data, "Финансы")
    assert len(back["sheetOrder"]) == 3


def test_shared_formulas_are_expanded_on_export() -> None:
    snapshot = finance_workbook("x")
    sheet = snapshot["sheets"][snapshot["sheetOrder"][0]]
    sheet["cellData"]["20"] = {
        "0": {"f": "=B21*2", "si": "abc", "v": 0},
        "1": {"si": "abc", "v": 0},
    }
    wb = load_workbook(io.BytesIO(snapshot_to_xlsx(snapshot)))
    assert wb["Обзор"]["B21"].value == "=C21*2"


def test_csv_export_neutralizes_formulas() -> None:
    snapshot = finance_workbook("x")
    sheet = snapshot["sheets"][snapshot["sheetOrder"][0]]
    sheet["cellData"]["0"]["0"] = {"v": '=HYPERLINK("http://evil")', "t": 1}
    text = snapshot_to_csv(snapshot).decode("utf-8-sig")
    assert text.splitlines()[0].startswith("\"'=HYPERLINK")


async def test_import_and_export_endpoints(client: httpx.AsyncClient) -> None:
    await register(client, "files@example.com")
    ws = await personal_workspace(client)

    r = await client.post(
        "/api/workbooks/import",
        data={"workspace_id": ws},
        files={"file": ("Отчёт.xlsx", _sample_xlsx(), "application/octet-stream")},
    )
    assert r.status_code == 201, r.text
    wb_id = r.json()["id"]
    assert r.json()["title"] == "Отчёт"

    r = await client.post(f"/api/workbooks/{wb_id}/export", params={"format": "xlsx"})
    assert r.status_code == 200
    assert (
        "filename*=UTF-8''%D0%9E%D1%82%D1%87%D1%91%D1%82.xlsx" in r.headers["content-disposition"]
    )
    assert r.headers["x-content-type-options"] == "nosniff"
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Данные", "Второй"]

    r = await client.post(
        "/api/workbooks/import",
        data={"workspace_id": ws},
        files={"file": ("data.csv", "имя;сумма\nа;1,5\nб;=1+1\n".encode(), "text/csv")},
    )
    assert r.status_code == 201
    wb = (await client.get(f"/api/workbooks/{r.json()['id']}")).json()
    cells = next(iter(wb["snapshot"]["sheets"].values()))["cellData"]
    assert cells["1"]["1"] == {"v": 1.5, "t": 2}
    assert cells["2"]["1"]["t"] == 4  # formulas from CSV stay text

    r = await client.post(
        "/api/workbooks/import",
        data={"workspace_id": ws},
        files={"file": ("bad.xlsx", b"not a zip", "application/octet-stream")},
    )
    assert r.status_code == 422
    r = await client.post(
        "/api/workbooks/import",
        data={"workspace_id": ws},
        files={"file": ("x.exe", b"MZ", "application/octet-stream")},
    )
    assert r.status_code == 415
